from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT


def export_review_to_docx(review, sections):
    document = Document()

    if "name" in sections:
        h = document.add_heading(review.title, level=1)
        h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p = document.add_paragraph(review.get_review_type_display())
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_paragraph("")

    if "authors" in sections:
        authors = list()
        authors.append(review.author.profile.get_screen_name())
        for author in review.co_authors.all():
            authors.append(author.profile.get_screen_name())
        p = document.add_paragraph(", ".join(authors))
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        document.add_paragraph("")

    if "description" in sections:
        if review.description:
            document.add_paragraph(review.description)

    document.add_heading("Planning", level=2)

    if review.objective:
        document.add_paragraph(review.objective)

    # PICOC
    if "picoc" in sections:
        document.add_heading("PICOC", level=3)

        p = document.add_paragraph("", style="List Bullet")
        p.add_run("Population: ").bold = True
        p.add_run(review.population)

        p = document.add_paragraph("", style="List Bullet")
        p.add_run("Intervention: ").bold = True
        p.add_run(review.intervention)

        p = document.add_paragraph("", style="List Bullet")
        p.add_run("Comparison: ").bold = True
        p.add_run(review.comparison)

        p = document.add_paragraph("", style="List Bullet")
        p.add_run("Outcome: ").bold = True
        p.add_run(review.outcome)

        p = document.add_paragraph("", style="List Bullet")
        p.add_run("Context: ").bold = True
        p.add_run(review.context)

    # Research Questions
    if "research_questions" in sections:
        document.add_heading("Research Questions", level=3)

        for question in review.research_questions.all():
            document.add_paragraph(question.question, style="List Number")

    # Keywords and Synonym
    if "keywords_synonyms" in sections:
        document.add_heading("Keywords and Synonyms", level=3)

        table = document.add_table(rows=1, cols=2)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Keyword"
        hdr_cells[1].text = "Synonyms"

        for keyword in review.get_keywords():
            row_cells = table.add_row().cells
            row_cells[0].text = keyword.description
            row_cells[1].text = ", ".join(keyword.synonyms.all().values_list("description", flat=True))

    # Search String
    if "search_string" in sections:
        document.add_heading("Search String", level=3)
        document.add_paragraph(review.get_generic_search_string().search_string)

    # Sources
    if "sources" in sections:
        document.add_heading("Sources", level=3)

        for source in review.sources.all():
            text = source.name
            if source.url:
                text = "{0} ({1})".format(source.name, source.url)
            document.add_paragraph(text, style="List Bullet")

    # Selection Criteria
    if "selection_criteria" in sections:
        document.add_heading("Selection Criteria", level=3)

        p = document.add_paragraph()
        p.add_run("Inclusion Criteria:").bold = True
        for criteria in review.get_inclusion_criterias():
            document.add_paragraph(criteria.description, style="List Bullet")

        p = document.add_paragraph()
        p.add_run("Exclusion Criteria:").bold = True
        for criteria in review.get_exclusion_criterias():
            document.add_paragraph(criteria.description, style="List Bullet")

    # Quality Assessment Checklist
    if "quality_assessment_checklist" in sections:
        document.add_heading("Quality Assessment Checklist", level=3)

        p = document.add_paragraph()
        p.add_run("Questions:").bold = True
        for quality_question in review.get_quality_assessment_questions():
            document.add_paragraph(quality_question.description, style="List Bullet")

        p = document.add_paragraph()
        p.add_run("Answers:").bold = True
        for quality_answer in review.get_quality_assessment_answers():
            document.add_paragraph(quality_answer.description, style="List Bullet")

    # Data Extraction Form
    if "data_extraction_form" in sections:
        document.add_heading("Data Extraction Form", level=3)
        for field in review.get_data_extraction_fields():
            document.add_paragraph(field.description, style="List Bullet")

    # Classification scheme (mapping studies)
    if "classification_scheme" in sections and review.facets.exists():
        document.add_heading("Classification Scheme", level=3)
        for facet in review.facets.prefetch_related("categories"):
            p = document.add_paragraph()
            p.add_run(f"{facet.name} ({facet.get_kind_display()}): ").bold = True
            p.add_run(", ".join(facet.categories.values_list("name", flat=True)))

    # Procedures and strategies recommended by the guidelines (K 5.4, P 5.1.5)
    for key, title in (
        ("selection_procedure", "Study Selection Procedure"),
        ("synthesis_strategy", "Synthesis Strategy"),
        ("validity_threats", "Validity Threats"),
    ):
        if key in sections and getattr(review, key):
            document.add_heading(title, level=3)
            document.add_paragraph(getattr(review, key))

    # Test set of known papers (P 5.1.2)
    if "test_set" in sections and review.test_set.exists():
        from parsifal.apps.reviews.services import mapping

        report = mapping.test_set_report(review)
        document.add_heading("Search Evaluation: Known Papers", level=3)
        document.add_paragraph(
            "{0} of {1} known papers were retrieved by the search ({2:.0%}).".format(
                report["found"], report["total"], report["recall"] or 0
            )
        )
        for item in report["items"]:
            status = "found" if item.matched_article_id else "not found"
            document.add_paragraph(f"{item.title} — {status}", style="List Bullet")

    # Conducting
    document.add_heading("Conducting", level=2)

    # Digital Libraries Search Strings
    if "source_search_strings" in sections:
        document.add_heading("Digital Libraries Search Strings", level=3)
        for search_session in review.get_latest_source_search_strings():
            p = document.add_paragraph()
            p.add_run("{0}:".format(search_session.source.name)).bold = True
            document.add_paragraph(search_session.search_string)
            document.add_paragraph()

    if "number_imported_studies" in sections:
        document.add_heading("Imported Studies", level=3)
        for source in review.sources.all():
            p = document.add_paragraph(style="List Bullet")
            p.add_run("{0}: ".format(source.name)).bold = True
            count = review.article_set.filter(source=source).count()
            p.add_run(str(count))

    # Search documentation (K 6.1.4 Table 2)
    if "search_documentation" in sections:
        sessions = review.get_latest_source_search_strings()
        if sessions:
            document.add_heading("Search Documentation", level=3)
            table = document.add_table(rows=1, cols=5)
            for cell, text in zip(table.rows[0].cells, ("Source", "Type", "Date", "Years", "Results")):
                cell.text = text
            for session in sessions:
                years = ""
                if session.years_from or session.years_to:
                    years = f"{session.years_from or ''}–{session.years_to or ''}"
                cells = table.add_row().cells
                cells[0].text = session.source.name
                cells[1].text = session.get_search_type_display()
                cells[2].text = session.searched_at.isoformat() if session.searched_at else ""
                cells[3].text = years
                cells[4].text = "" if session.result_count is None else str(session.result_count)

    if "quality_assessment" in sections:
        from parsifal.apps.reviews.services import quality

        document.add_heading("Quality Assessment", level=3)
        summary = quality.summary(review)
        if summary["has_checklist"]:
            document.add_paragraph("Maximum score: {0}. Cutoff: {1}.".format(summary["max_score"], summary["cutoff"]))
            table = document.add_table(rows=1, cols=3)
            for cell, text in zip(table.rows[0].cells, ("Study", "Score", "Final selection")):
                cell.text = text
            for row in summary["articles"]:
                cells = table.add_row().cells
                cells[0].text = row["title"] or ""
                cells[1].text = str(row["score"])
                cells[2].text = "yes" if row["passes_cutoff"] else "no"
        else:
            document.add_paragraph("No quality assessment checklist was defined for this review.")

    if "data_extraction" in sections:
        from parsifal.apps.reviews.services import reports

        document.add_heading("Data Extraction", level=3)
        columns, rows = reports.extraction_rows(review)
        if rows and len(columns) > 1:
            table = document.add_table(rows=1, cols=len(columns))
            for cell, text in zip(table.rows[0].cells, columns):
                cell.text = text
            for row in rows:
                cells = table.add_row().cells
                for cell, column in zip(cells, columns):
                    cell.text = str(row[column] or "")
        else:
            document.add_paragraph("No data has been extracted yet.")

    if "classification_results" in sections and review.facets.exists():
        from parsifal.apps.reviews.services import classification

        document.add_heading("Classification Results", level=3)
        for facet in review.facets.all():
            frequencies = classification.facet_frequencies(review=review, facet_id=facet.pk)
            p = document.add_paragraph()
            p.add_run(f"{facet.name}: ").bold = True
            p.add_run(
                "; ".join(f"{row['category']} ({row['count']})" for row in frequencies["categories"])
                or "no categories"
            )

    if "data_analysis" in sections:
        from parsifal.apps.reviews.services import analytics

        document.add_heading("Data Analysis", level=3)
        p = document.add_paragraph()
        p.add_run("Studies per source (imported / accepted): ").bold = True
        p.add_run(
            "; ".join(
                f"{row['source']} {row['total']} / {row['accepted']}" for row in analytics.selection_by_source(review)
            )
            or "none"
        )
        p = document.add_paragraph()
        p.add_run("Final selection per year: ").bold = True
        p.add_run(
            "; ".join(
                f"{row['year'] if row['year'] is not None else 'unknown'}: {row['count']}"
                for row in analytics.final_selection_by_year(review)
            )
            or "none"
        )

    if "process_checklist" in sections:
        from parsifal.apps.reviews.services import mapping

        checklist = mapping.checklist(review)
        if checklist["applied"]:
            document.add_heading("Process Evaluation (Petersen et al. 2015)", level=3)
            document.add_paragraph(
                "{0} of {1} recommended actions applied ({2:.0%}).".format(
                    checklist["applied"], checklist["total"], checklist["ratio"]
                )
            )
            for row in checklist["items"]:
                if row["applied"]:
                    text = f"[{row['phase']}] {row['label']}"
                    if row["notes"]:
                        text += f" — {row['notes']}"
                    document.add_paragraph(text, style="List Bullet")

    return document
