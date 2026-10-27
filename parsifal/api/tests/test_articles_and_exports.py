import io
import json
from pathlib import Path

from openpyxl import load_workbook

from parsifal.api.tests.base import API, ApiTestCase
from parsifal.apps.reviews.models import Article
from parsifal.apps.reviews.tests.factories import ArticleFactory, SelectionCriteriaFactory, SourceFactory

DATA = Path(__file__).resolve().parents[2] / "apps" / "reviews" / "tests" / "data"


class TestArticlesApi(ApiTestCase):
    def test_import_json_and_file(self):
        response = self.api(
            "post",
            self.review_path("/articles/import"),
            {"source_id": self.source.pk, "bibtex": (DATA / "scopus_sample.bib").read_text()},
        )
        self.assertEqual(201, response.status_code)
        body = response.json()
        self.assertEqual((2, 0, []), (body["imported"], body["failed"], body["warnings"]))
        self.assertEqual(2, len(body["article_ids"]))

        with open(DATA / "ieee_sample.bib", "rb") as handle:
            response = self.client.post(
                f"{API}{self.review_path('/articles/import/file')}", {"source_id": self.source.pk, "file": handle}
            )
        self.assertEqual(201, response.status_code)
        self.assertEqual(2, response.json()["imported"])
        self.assertEqual(4, self.review.article_set.count())

        with self.subTest(msg="empty bibtex is a validation error"):
            response = self.api(
                "post", self.review_path("/articles/import"), {"source_id": self.source.pk, "bibtex": " "}
            )
            self.assertEqual(422, response.status_code)
            self.assertEqual("validation_error", response.json()["code"])

    def test_list_pagination_and_filters(self):
        ArticleFactory.create_batch(3, review=self.review, source=self.source)
        ArticleFactory(review=self.review, source=self.source, status="A", title="Chosen")
        body = self.api("get", self.review_path("/articles?limit=2")).json()
        self.assertEqual((4, 2), (body["count"], len(body["items"])))
        body = self.api("get", self.review_path("/articles?status=A")).json()
        self.assertEqual(["Chosen"], [a["title"] for a in body["items"]])
        self.assertEqual(422, self.api("get", self.review_path("/articles?status=Z")).status_code)
        ArticleFactory(review=self.review, source=self.source, title="By DOI", doi="10.1000/xyz.42")
        ArticleFactory(review=self.review, source=self.source, title="By key", bibtex_key="Doe2024Mapping")
        for query, titles in (("chosen", ["Chosen"]), ("doi.org/10.1000/XYZ", ["By DOI"]), ("doe2024", ["By key"])):
            with self.subTest(q=query):
                items = self.api("get", self.review_path("/articles"), QUERY_STRING=f"q={query}").json()["items"]
                self.assertEqual(titles, [a["title"] for a in items])

    def test_create_get_update_delete(self):
        criteria = SelectionCriteriaFactory(review=self.review)
        response = self.api(
            "post", self.review_path("/articles"), {"source_id": self.source.pk, "title": "Manual", "year": "2023"}
        )
        self.assertEqual(201, response.status_code)
        article = response.json()
        self.assertEqual((2023, self.author.username), (article["year_int"], article["created_by"]))

        response = self.api(
            "patch",
            self.review_path(f"/articles/{article['id']}"),
            {"status": "A", "selection_criteria_id": criteria.pk, "comments": "fits"},
        )
        self.assertEqual(
            ("A", criteria.pk, "fits"),
            tuple(response.json()[k] for k in ("status", "selection_criteria_id", "comments")),
        )
        self.assertEqual(article["id"], self.api("get", self.review_path(f"/articles/{article['id']}")).json()["id"])
        self.assertEqual(204, self.api("delete", self.review_path(f"/articles/{article['id']}")).status_code)
        self.assertEqual(404, self.api("get", self.review_path(f"/articles/{article['id']}")).status_code)

    def test_bulk_status_and_scoping(self):
        a1, a2 = ArticleFactory.create_batch(2, review=self.review, source=self.source)
        foreign = ArticleFactory()
        response = self.api(
            "post", self.review_path("/articles/status"), {"ids": [a1.pk, a2.pk], "status": "R", "comments": "no"}
        )
        self.assertEqual(2, response.json()["count"])
        self.assertEqual({"R"}, set(Article.objects.filter(pk__in=[a1.pk, a2.pk]).values_list("status", flat=True)))
        response = self.api("post", self.review_path("/articles/status"), {"ids": [foreign.pk], "status": "A"})
        self.assertEqual(404, response.status_code)
        foreign.refresh_from_db()
        self.assertEqual("U", foreign.status)

    def test_duplicates(self):
        a1 = ArticleFactory(review=self.review, source=self.source, title="Dup Title")
        a2 = ArticleFactory(review=self.review, source=self.source, title="dup title")
        body = self.api("get", self.review_path("/articles/duplicates")).json()
        self.assertEqual([[a1.pk, a2.pk]], [[a["id"] for a in group] for group in body["groups"]])
        marked = self.api("post", self.review_path("/articles/duplicates/resolve"), {}).json()
        self.assertEqual([a2.pk], [a["id"] for a in marked])
        self.assertEqual("D", marked[0]["status"])

    def test_bulk_delete_and_clear_source(self):
        other = SourceFactory()
        self.review.sources.add(other)
        mine = [ArticleFactory(review=self.review, source=self.source).pk for _ in range(3)]
        other_source = ArticleFactory(review=self.review, source=other)
        foreign = ArticleFactory()
        with self.subTest(msg="one foreign id means nothing is deleted"):
            response = self.api("post", self.review_path("/articles/delete"), {"ids": [mine[0], foreign.pk]})
            self.assertEqual(404, response.status_code)
            self.assertEqual(3, self.review.article_set.filter(source=self.source).count())
        with self.subTest(msg="delete by ids"):
            response = self.api("post", self.review_path("/articles/delete"), {"ids": mine[:2]})
            self.assertEqual((200, 2), (response.status_code, response.json()["count"]))
        with self.subTest(msg="clear a source"):
            response = self.api("post", self.review_path("/articles/clear"), {"source_id": self.source.pk})
            self.assertEqual((200, 1), (response.status_code, response.json()["count"]))
            self.assertEqual([other_source.pk], list(self.review.article_set.values_list("pk", flat=True)))
        with self.subTest(msg="source of another review, bad status, outsider"):
            self.assertEqual(
                404, self.api("post", self.review_path("/articles/clear"), {"source_id": 99999}).status_code
            )
            self.assertEqual(
                422,
                self.api(
                    "post", self.review_path("/articles/clear"), {"source_id": other.pk, "status": "Z"}
                ).status_code,
            )
            outsider = self.client_for(self.outsider)
            self.assertEqual(
                404,
                self.api(
                    "post", self.review_path("/articles/clear"), {"source_id": other.pk}, client=outsider
                ).status_code,
            )
            self.assertEqual(1, self.review.article_set.count())

    def test_duplicates_report_reasons_and_conflicts(self):
        a = ArticleFactory(review=self.review, source=self.source, title="Using CNN-SSD", doi="10.1/a")
        b = ArticleFactory(review=self.review, source=self.source, title="Using CNN\u2013SSD", doi="10.1/A")
        c = ArticleFactory(review=self.review, source=self.source, title="Twin", doi="10.2/x")
        d = ArticleFactory(review=self.review, source=self.source, title="Twin", doi="10.2/y")
        body = self.api("get", self.review_path("/articles/duplicates")).json()
        self.assertEqual([[a.pk, b.pk]], [[x["id"] for x in group] for group in body["groups"]])
        self.assertEqual(
            [{"article_ids": [a.pk, b.pk], "match": "doi+title", "needs_review": False, "reason": ""}], body["details"]
        )
        self.assertEqual([[c.pk, d.pk]], [conflict["article_ids"] for conflict in body["conflicts"]])


class TestAnalyticsAndExports(ApiTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        ArticleFactory(review=cls.review, source=cls.source, status="A", year="2020", title="One")
        ArticleFactory(review=cls.review, source=cls.source, status="R", year="2021", title="Two")

    def test_analytics(self):
        rows = self.api("get", self.review_path("/analytics/selection-by-source")).json()
        self.assertEqual([(self.source.name, 2, 1)], [(r["source"], r["total"], r["accepted"]) for r in rows])
        rows = self.api("get", self.review_path("/analytics/final-by-year")).json()
        self.assertEqual([{"year": 2020, "count": 1}], rows)
        summary = self.api("get", self.review_path("/analytics/selection-summary")).json()
        self.assertEqual(
            (2, 1, 1, 0.5), (summary["total"], summary["accepted"], summary["rejected"], summary["accepted_share"])
        )
        usage = self.api("get", self.review_path("/analytics/criteria-usage")).json()
        self.assertEqual((1, 1), (usage["rejected_without_criterion"], usage["accepted_without_criterion"]))
        years = self.api("get", self.review_path("/analytics/studies-by-year")).json()
        self.assertEqual([(2020, 1, 0), (2021, 0, 1)], [(r["year"], r["accepted"], r["rejected"]) for r in years])
        overview = self.api("get", self.review_path("/analytics/quality")).json()
        self.assertEqual(
            (1, 0, 0, 1, None),
            tuple(overview[key] for key in ("accepted", "questions", "assessed", "not_assessed", "cutoff")),
        )
        choices = self.api("get", self.review_path("/analytics/extraction")).json()
        self.assertEqual({"studies": 1, "done": 0, "fields": []}, choices)
        self.assertEqual(
            404,
            self.api("get", self.review_path("/analytics/quality"), client=self.client_for(self.outsider)).status_code,
        )

    def test_exports(self):
        with self.subTest(msg="json"):
            response = self.api("get", self.review_path("/export/articles.json"))
            self.assertEqual(200, response.status_code)
            self.assertEqual({"One", "Two"}, {row["title"] for row in json.loads(response.content)})
        with self.subTest(msg="csv"):
            response = self.api("get", self.review_path("/export/articles.csv"))
            self.assertTrue(response["Content-Type"].startswith("text/csv"))
            self.assertIn("bibtex_key,title", response.content.decode())
        with self.subTest(msg="xlsx"):
            response = self.api("get", self.review_path("/export/extraction.xlsx"))
            sheet = load_workbook(io.BytesIO(response.content)).active
            self.assertEqual("article", sheet.cell(row=1, column=1).value)
            self.assertEqual("One", sheet.cell(row=2, column=1).value)
        with self.subTest(msg="bad format"):
            self.assertEqual(422, self.api("get", self.review_path("/export/articles.pdf")).status_code)
        with self.subTest(msg="docx"):
            response = self.api("get", self.review_path("/export/report.docx?sections=name,sources"))
            self.assertEqual(200, response.status_code)
            self.assertIn(f'filename="{self.review.name}.docx"', response["Content-Disposition"])
            self.assertEqual(b"PK", response.content[:2])
